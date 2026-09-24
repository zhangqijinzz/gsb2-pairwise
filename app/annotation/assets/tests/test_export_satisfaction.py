import json
import importlib.util
import hashlib
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
from zipfile import ZipFile

try:
    import openpyxl
except ImportError:  # pragma: no cover - the shipped exporter does not need it
    openpyxl = None


ASSETS = Path(__file__).resolve().parents[1]
EXPORTER = ASSETS / "export_satisfaction.py"
CHECKER = ASSETS / "check_submission.py"
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


class ExportSatisfactionTests(unittest.TestCase):
    def test_repair_consistency_blocks_contradictory_saved_results(self):
        spec = importlib.util.spec_from_file_location("repair_exporter", EXPORTER)
        exporter = importlib.util.module_from_spec(spec)
        sys.path.insert(0, str(ASSETS))
        try:
            spec.loader.exec_module(exporter)
        finally:
            sys.path.pop(0)
        e = self.evaluation("h")
        e["issues"] = [{"kind": "bug", "description": "空值提交无提示", "evidence": "code/form.ts"}]
        self.assertTrue(exporter._validate_evaluation(e, "round"))
        e["scores"][0] = 4
        self.assertTrue(exporter._validate_evaluation(e, "round"))
        e["nextPrompt"] = "补充错误提示"
        e["nextPromptType"] = "Bug修复"
        self.assertTrue(exporter._validate_evaluation(e, "round"))
        e["nextPrompt"] = "修复空值提交无提示的问题，显示错误信息"
        self.assertEqual(exporter._validate_evaluation(e, "round"), [])
        e["issues"] = []
        self.assertTrue(exporter._validate_evaluation(e, "round"))
        e["nextPrompt"] = ""
        e["nextPromptType"] = ""
        e["scores"] = [5, 5, 4, 5, 4]
        e["issues"] = [{"kind": "process", "description": "重复检索", "evidence": "原始调用记录"}]
        self.assertEqual(exporter._validate_evaluation(e, "round"), [])
        e["scores"] = [5, 5, 5, 5, 5]
        e["nextPrompt"] = "修复重复检索"
        e["nextPromptType"] = "Bug修复"
        e["scores"][0] = 5
        self.assertTrue(exporter._validate_evaluation(e, "round"))

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def capture(self, name):
        source = self.root / "sources" / name
        (source / "trace" / "subagents").mkdir(parents=True)
        (source / "trace" / "session.jsonl").write_text('{"type":"user"}\n', encoding="utf-8")
        (source / "trace" / "subagents" / "child.jsonl").write_text("{}\n", encoding="utf-8")
        (source / "code").mkdir()
        (source / "code" / "result.txt").write_text("model artifact\n", encoding="utf-8")
        return {
            "id": name,
            "dir": str(source),
            "tracePath": "trace/session.jsonl",
            "codePath": "code",
            "hash": "capture-" + name,
            "createdAt": 1,
        }

    def evaluation(self, evidence_hash, *, scores=None, status="ready", harness="2.1.197"):
        review = self.root / "reviews" / ("review-" + evidence_hash)
        (review / "verification").mkdir(parents=True, exist_ok=True)
        (review / "verification" / "check.txt").write_text("independent check\n", encoding="utf-8")
        (review / "input.json").write_text("{}\n", encoding="utf-8")
        (review / "evaluation.json").write_text("{}\n", encoding="utf-8")
        return {
            "id": "eval-" + evidence_hash,
            "createdAt": 2,
            "skillHash": "skill-1",
            "model": "review-model",
            "evidenceHash": evidence_hash,
            "status": status,
            "scores": scores if scores is not None else [5, 4, 4, 4, 4],
            "descriptions": ["交付依据", "遵循依据", "规划依据", "推理依据", "执行依据"],
            "taskType": "feature迭代",
            "difficulty": "中等",
            "language": "Python",
            "environment": "已容器化，可一键起环境",
            "harnessVersion": harness,
            "os": "MacOS/Linux",
            "evidence": ["verification/check.txt:1"],
            "missing": [],
            "issues": [],
            "nextPrompt": "",
            "nextPromptType": "",
            "reviewPath": str(review),
            "reviewHash": "review-hash-" + evidence_hash,
        }

    def round(self, prompt_id, capture_id, order, *, prompt="Do the work", status="complete", evidence_hash=None, evaluations=None, session_id="session-1", reason=""):
        evidence_hash = evidence_hash or "evidence-" + prompt_id
        if evaluations is None:
            evaluations = [self.evaluation(evidence_hash)]
        return {
            "promptId": prompt_id,
            "sessionId": session_id,
            "prompt": prompt,
            "order": order,
            "status": status,
            "reason": reason,
            "evidenceHash": evidence_hash,
            "version": "round-v1",
            "cwd": "/workspace/project",
            "captureId": capture_id,
            "evaluations": evaluations,
        }

    def case(self, task_id, rounds, captures, *, completed=True):
        sha = (task_id[-1:] or "a") * 40
        if not all(ch in "0123456789abcdef" for ch in sha):
            sha = "a" * 40
        return {
            "taskId": task_id,
            "projectId": "project-1",
            "taskName": "Task " + task_id,
            "initialSha": sha,
            "snapshotUrl": "https://github.com/example/repo/commit/" + sha,
            "completed": completed,
            "rounds": rounds,
            "captures": captures,
        }

    @staticmethod
    def attachment_root(output, task_id, capture_id):
        manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
        case = next(item for item in manifest["cases"] if item["taskId"] == task_id)
        capture = next(item for item in case["captures"] if item["id"] == capture_id)
        return output / capture["attachmentRoot"]

    def payload(self, cases, *, issues=None, submitted_at="2026-09-12T10:30:00+08:00"):
        return {
            "projectName": "Project Demo",
            "submitter": "Reviewer",
            "submittedAt": submitted_at,
            "cases": cases,
            "issues": issues or [],
        }

    def run_export(self, payload, *, draft=False, output_name=None):
        input_path = self.root / "input.json"
        input_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        output = self.root / (output_name or ("draft" if draft else "formal"))
        command = [sys.executable, str(EXPORTER), "--input", str(input_path), "--output", str(output)]
        if draft:
            command.append("--draft")
        proc = subprocess.run(command, text=True, capture_output=True)
        try:
            result = json.loads(proc.stdout)
        except json.JSONDecodeError as exc:
            self.fail(f"exporter stdout is not JSON: {proc.stdout!r}; stderr={proc.stderr!r}; {exc}")
        return proc, result, output

    @unittest.skipIf(openpyxl is None, "openpyxl is only needed for workbook behavior tests")
    def test_multi_case_keeps_earlier_failure_and_uses_last_matching_evaluation(self):
        capture_a = self.capture("capture-a")
        capture_b = self.capture("capture-b")
        stale = self.evaluation("old-hash", scores=[5, 5, 5, 5, 5])
        first_match = self.evaluation("hash-a", scores=[4, 4, 4, 4, 4])
        latest_match = self.evaluation("hash-a", scores=[3, 3, 4, 5, 3])
        first = self.round("prompt-a", "capture-a", 1, prompt="first failed attempt", evidence_hash="hash-a", evaluations=[stale, first_match, latest_match])
        second = self.round("prompt-b", "capture-b", 2, prompt="later success", session_id="session-2")
        cases = [self.case("task-b", [second], [capture_b]), self.case("task-a", [first], [capture_a])]

        proc, result, output = self.run_export(self.payload(cases))

        self.assertEqual(proc.returncode, 0, proc.stderr)
        self.assertEqual(result, {
            "outputPath": str(output.resolve() / "submission.xlsx"),
            "reportPath": str(output.resolve() / "report.md"),
            "rows": 2,
            "issues": [],
        })
        workbook = openpyxl.load_workbook(result["outputPath"], data_only=False)
        sheet = workbook.active
        self.assertEqual([sheet.cell(row, 1).value for row in (2, 3)], ["first failed attempt", "later success"])
        self.assertEqual([sheet.cell(2, col).value for col in (13, 15, 17, 19, 21)], [3, 3, 4, 5, 3])
        copied_capture = self.attachment_root(output, "task-a", "capture-a")
        self.assertTrue((copied_capture / "trace" / "subagents" / "child.jsonl").is_file())
        self.assertEqual(sheet["E2"].value, str((copied_capture / "trace" / "session.jsonl").resolve()))
        self.assertTrue(Path(sheet["E2"].value).is_file())
        manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
        self.assertEqual([r["prompt"] for r in manifest["rounds"]], ["first failed attempt", "later success"])
        self.assertEqual(manifest["rounds"][0]["evaluationId"], latest_match["id"])

    def test_duplicate_session_and_prompt_identity_blocks_formal_export(self):
        capture_a = self.capture("capture-a")
        capture_b = self.capture("capture-b")
        duplicate_a = self.round("same", "capture-a", 1, session_id="same-session")
        duplicate_b = self.round("same", "capture-b", 1, session_id="same-session")

        proc, result, output = self.run_export(self.payload([
            self.case("task-a", [duplicate_a], [capture_a]),
            self.case("task-b", [duplicate_b], [capture_b]),
        ]))

        self.assertNotEqual(proc.returncode, 0)
        self.assertEqual(result["outputPath"], "")
        self.assertEqual(result["rows"], 0)
        self.assertTrue(any("duplicate SessionID + PromptID" in issue for issue in result["issues"]))
        self.assertFalse((output / "submission.xlsx").exists())

    @unittest.skipIf(openpyxl is None, "openpyxl is needed for workbook checks")
    def test_export_uses_card_type_for_all_rounds_over_cached_ai_classification(self):
        for task_type in ("0-1代码生成", "Feature迭代", "Bug修复", "代码理解", "工程化", "代码测试", "代码重构"):
            with self.subTest(task_type=task_type):
                capture = self.capture("capture-" + task_type)
                rounds = [self.round("first-" + task_type, capture["id"], 1), self.round("second-" + task_type, capture["id"], 2)]
                case = self.case("task-a", rounds, [capture])
                case["taskType"] = task_type
                proc, result, _ = self.run_export(self.payload([case]), output_name=task_type)
                self.assertEqual(proc.returncode, 0, proc.stderr + str(result))
                sheet = openpyxl.load_workbook(result["outputPath"]).active
                self.assertEqual([sheet["J2"].value, sheet["J3"].value], [task_type, task_type])
                checked = subprocess.run([sys.executable, str(CHECKER), result["outputPath"]], capture_output=True, text=True)
                self.assertEqual(checked.returncode, 0, checked.stdout + checked.stderr)

    def test_draft_keeps_conflicting_duplicate_rounds_and_flags_identity(self):
        capture = self.capture("capture-a")
        first = self.round("same", "capture-a", 1, prompt="first text", status="conflict", evaluations=[])
        second = self.round("same", "capture-a", 2, prompt="different text", status="pending", evaluations=[])
        case = self.case("task-a", [first, second], [capture], completed=False)

        proc, result, output = self.run_export(self.payload([case]), draft=True)

        self.assertEqual(proc.returncode, 0, proc.stderr)
        self.assertEqual(result["rows"], 2)
        self.assertTrue(any("duplicate SessionID + PromptID" in issue for issue in result["issues"]))
        manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
        self.assertEqual([item["prompt"] for item in manifest["rounds"]], ["first text", "different text"])

    @unittest.skipIf(openpyxl is None, "openpyxl is only needed for workbook behavior tests")
    def test_formula_like_prompt_is_stored_as_literal_text(self):
        capture = self.capture("capture-a")
        prompt = "=HYPERLINK(\"https://example.invalid\",\"click\")"
        proc, result, _ = self.run_export(self.payload([self.case("task-a", [self.round("prompt-a", "capture-a", 1, prompt=prompt)], [capture])]))

        self.assertEqual(proc.returncode, 0, proc.stderr)
        sheet = openpyxl.load_workbook(result["outputPath"], data_only=False).active
        self.assertEqual(sheet["A2"].value, prompt)
        self.assertEqual(sheet["A2"].data_type, "s")

    @unittest.skipIf(openpyxl is None, "openpyxl is only needed for workbook behavior tests")
    def test_missing_trace_blocks_formal_but_draft_keeps_row_without_fake_values(self):
        capture = self.capture("capture-a")
        Path(capture["dir"], capture["tracePath"]).unlink()
        pending = self.round("prompt-a", "capture-a", 1, status="pending", evaluations=[])
        case = self.case("task-a", [pending], [capture], completed=False)
        payload = self.payload([case], submitted_at="")

        formal_proc, formal_result, _ = self.run_export(payload)
        draft_proc, draft_result, _ = self.run_export(payload, draft=True)

        self.assertNotEqual(formal_proc.returncode, 0)
        self.assertTrue(any("trace attachment is missing" in issue for issue in formal_result["issues"]))
        self.assertEqual(draft_proc.returncode, 0, draft_proc.stderr)
        self.assertEqual(draft_result["rows"], 1)
        sheet = openpyxl.load_workbook(draft_result["outputPath"], data_only=False).active
        self.assertEqual(sheet["A2"].value, "Do the work")
        for cell in ("E2", "F2", "G2", "H2", "I2", "J2", "K2", "L2", "M2", "N2", "Y2", "Z2"):
            self.assertIsNone(sheet[cell].value, cell)
        self.assertEqual(sheet["X2"].value, "Reviewer")
        self.assertTrue(any("trace attachment is missing" in issue for issue in draft_result["issues"]))

    def test_overlong_prompt_is_rejected_formally_and_preserved_as_draft_attachment(self):
        capture = self.capture("capture-a")
        prompt = "x" * 32768
        case = self.case("task-a", [self.round("prompt-a", "capture-a", 1, prompt=prompt)], [capture])

        formal_proc, formal_result, _ = self.run_export(self.payload([case]))
        draft_proc, draft_result, draft_output = self.run_export(self.payload([case]), draft=True)

        self.assertNotEqual(formal_proc.returncode, 0)
        self.assertTrue(any("32767" in issue for issue in formal_result["issues"]))
        self.assertEqual(draft_proc.returncode, 0, draft_proc.stderr)
        prompt_files = list((draft_output / "attachments").glob("*/original-prompts/*.txt"))
        self.assertEqual(len(prompt_files), 1)
        self.assertEqual(prompt_files[0].read_text(encoding="utf-8"), prompt)
        manifest = json.loads((draft_output / "manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["rounds"][0]["promptAttachment"], prompt_files[0].relative_to(draft_output).as_posix())

    def test_validation_ranges_cover_batches_larger_than_template_default(self):
        capture = self.capture("capture-a")
        rounds = [self.round(f"prompt-{i:03d}", "capture-a", i, session_id=f"session-{i // 10:03d}") for i in range(1, 121)]
        proc, result, _ = self.run_export(self.payload([self.case("task-a", rounds, [capture])]))

        self.assertEqual(proc.returncode, 0, proc.stderr)
        with ZipFile(result["outputPath"]) as archive:
            root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        refs = {validation.get("sqref") for validation in root.findall(NS + "dataValidations/" + NS + "dataValidation")}
        self.assertTrue({"F2:F121", "G2:G121", "M2:M121", "U2:U121", "Z2:Z121"}.issubset(refs))
        self.assertEqual(root.find(NS + "dimension").get("ref"), "A1:AB121")

    def test_excluded_round_is_manifested_with_reason_and_has_no_row(self):
        capture = self.capture("capture-a")
        excluded = self.round("excluded", "capture-a", 1, status="excluded", reason="tool result, not a human prompt", evaluations=[])
        valid = self.round("valid", "capture-a", 2)

        proc, result, output = self.run_export(self.payload([self.case("task-a", [excluded, valid], [capture])]))

        self.assertEqual(proc.returncode, 0, proc.stderr)
        self.assertEqual(result["rows"], 1)
        manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
        self.assertEqual([(r["promptId"], r["included"], r["reason"]) for r in manifest["rounds"]], [
            ("excluded", False, "tool result, not a human prompt"),
            ("valid", True, ""),
        ])

    def test_score_total_above_twenty_one_is_manifested_but_not_exported(self):
        capture = self.capture("capture-a")
        score_21 = self.round("score-21", "capture-a", 1, evaluations=[self.evaluation("evidence-score-21", scores=[5, 4, 4, 4, 4])], evidence_hash="evidence-score-21")
        score_22 = self.round("score-22", "capture-a", 2, evaluations=[self.evaluation("evidence-score-22", scores=[5, 5, 4, 4, 4])], evidence_hash="evidence-score-22")

        proc, result, output = self.run_export(self.payload([self.case("task-a", [score_21, score_22], [capture])]))

        self.assertEqual(proc.returncode, 0, proc.stderr)
        self.assertEqual(result["rows"], 1)
        manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
        self.assertEqual([(item["promptId"], item["included"], item["scoreTotal"]) for item in manifest["rounds"]], [
            ("score-21", True, 21),
            ("score-22", False, 22),
        ])
        self.assertIn("超过 21", manifest["rounds"][1]["reason"])

    def test_input_issues_block_formal_export_but_remain_in_draft(self):
        capture = self.capture("capture-a")
        case = self.case("task-a", [self.round("prompt-a", "capture-a", 1)], [capture])
        payload = self.payload([case], issues=["parser found an ambiguous event boundary"])

        formal_proc, formal_result, _ = self.run_export(payload)
        draft_proc, draft_result, _ = self.run_export(payload, draft=True)

        self.assertNotEqual(formal_proc.returncode, 0)
        self.assertTrue(any("parser found an ambiguous event boundary" in issue for issue in formal_result["issues"]))
        self.assertEqual(draft_proc.returncode, 0, draft_proc.stderr)
        self.assertTrue(any("parser found an ambiguous event boundary" in issue for issue in draft_result["issues"]))

    def test_checker_reads_identifiers_and_reports_validation_scope(self):
        capture = self.capture("capture-a")
        case = self.case("task-a", [self.round("prompt-a", "capture-a", 1)], [capture])
        proc, result, output = self.run_export(self.payload([case]))
        self.assertEqual(proc.returncode, 0, proc.stderr)

        checked = subprocess.run([
            sys.executable,
            str(CHECKER),
            result["outputPath"],
            "--rounds-json",
            str(output / "manifest.json"),
        ], text=True, capture_output=True)
        self.assertEqual(checked.returncode, 0, checked.stderr)
        report = json.loads(checked.stdout)
        self.assertEqual(report["field_status"], "PASS")
        self.assertEqual(report["rows_checked"], 1)
        self.assertEqual(report["round_identity_check"], "compared_to_supplied_manifest")
        self.assertEqual(report["validation_ranges"]["M"], "M2:M100")
        self.assertIn("score/reason consistency and natural prose", report["not_checked"])

    @unittest.skipIf(openpyxl is None, "openpyxl is only needed for workbook behavior tests")
    def test_checker_rejects_a_row_whose_score_total_exceeds_twenty_one(self):
        capture = self.capture("capture-a")
        case = self.case("task-a", [self.round("prompt-a", "capture-a", 1)], [capture])
        proc, result, output = self.run_export(self.payload([case]))
        self.assertEqual(proc.returncode, 0, proc.stderr)
        workbook = openpyxl.load_workbook(result["outputPath"])
        for column in (13, 15, 17, 19, 21):
            workbook.active.cell(2, column).value = 5
        workbook.save(result["outputPath"])

        checked = subprocess.run([
            sys.executable,
            str(CHECKER),
            result["outputPath"],
            "--rounds-json",
            str(output / "manifest.json"),
        ], text=True, capture_output=True)

        self.assertNotEqual(checked.returncode, 0)
        self.assertIn("score total exceeds 21", checked.stdout)

    @unittest.skipIf(openpyxl is None, "openpyxl is only needed for workbook behavior tests")
    def test_checker_rejects_a_dimension_score_below_three(self):
        capture = self.capture("capture-a")
        case = self.case("task-a", [self.round("prompt-a", "capture-a", 1)], [capture])
        proc, result, output = self.run_export(self.payload([case]))
        self.assertEqual(proc.returncode, 0, proc.stderr)
        workbook = openpyxl.load_workbook(result["outputPath"])
        workbook.active.cell(2, 13).value = 2
        workbook.save(result["outputPath"])

        checked = subprocess.run([
            sys.executable,
            str(CHECKER),
            result["outputPath"],
            "--rounds-json",
            str(output / "manifest.json"),
        ], text=True, capture_output=True)

        self.assertNotEqual(checked.returncode, 0)
        self.assertIn("numeric integer 3–5", checked.stdout)

    def test_same_frozen_input_preserves_values_except_absolute_export_location(self):
        capture = self.capture("capture-a")
        case = self.case("task-a", [self.round("prompt-a", "capture-a", 1)], [capture])
        payload = self.payload([case])

        first_proc, first, _ = self.run_export(payload)
        input_path = self.root / "input.json"
        second_output = self.root / "formal-second"
        second_proc = subprocess.run([
            sys.executable,
            str(EXPORTER),
            "--input", str(input_path),
            "--output", str(second_output),
        ], text=True, capture_output=True)
        second = json.loads(second_proc.stdout)

        self.assertEqual(first_proc.returncode, 0, first_proc.stderr)
        self.assertEqual(second_proc.returncode, 0, second_proc.stderr)
        first_rows = list(openpyxl.load_workbook(first["outputPath"]).active.values)
        second_rows = list(openpyxl.load_workbook(second["outputPath"]).active.values)
        self.assertEqual(len(first_rows), len(second_rows))
        for index, (left, right) in enumerate(zip(first_rows, second_rows)):
            if index == 0:
                self.assertEqual(left, right)
                continue
            self.assertEqual(left[:4] + left[5:], right[:4] + right[5:])
            self.assertTrue(Path(left[4]).is_absolute())
            self.assertTrue(Path(right[4]).is_absolute())
            self.assertEqual(Path(left[4]).read_bytes(), Path(right[4]).read_bytes())

    def test_checker_rejects_relative_trace_even_when_file_exists(self):
        capture = self.capture("capture-a")
        proc, result, output = self.run_export(self.payload([
            self.case("task-a", [self.round("prompt-a", "capture-a", 1)], [capture]),
        ]))
        self.assertEqual(proc.returncode, 0, proc.stderr)
        workbook = openpyxl.load_workbook(result["outputPath"])
        workbook.active["E2"] = os.path.relpath(workbook.active["E2"].value, output)
        workbook.save(result["outputPath"])
        checked = subprocess.run([sys.executable, str(CHECKER), result["outputPath"]], text=True, capture_output=True)
        self.assertNotEqual(checked.returncode, 0)
        self.assertIn("local absolute JSONL file path", checked.stdout)

    def test_combined_export_separates_tasks_without_counting_blank_rows(self):
        a, b = self.capture("a"), self.capture("b")
        payload = self.payload([
            self.case("task-a", [self.round("p1", "a", 1), self.round("p2", "a", 2)], [a]),
            self.case("task-b", [self.round("p3", "b", 1)], [b]),
        ])
        payload["separateTasks"] = True
        proc, result, _ = self.run_export(payload)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        sheet = openpyxl.load_workbook(result["outputPath"]).active
        self.assertEqual(result["rows"], 3)
        self.assertEqual(sheet.max_row, 5)
        self.assertEqual([sheet.cell(row, 3).value for row in (2, 3, 5)], ["p1", "p2", "p3"])
        self.assertTrue(all(sheet.cell(4, col).value is None for col in range(1, 29)))

    def test_capture_symlink_is_rejected_without_copying_outside_content(self):
        capture = self.capture("capture-a")
        outside = self.root / "outside-secret.txt"
        outside.write_text("must not copy", encoding="utf-8")
        os.symlink(outside, Path(capture["dir"]) / "trace" / "outside-link")
        case = self.case("task-a", [self.round("prompt-a", "capture-a", 1)], [capture])

        proc, result, output = self.run_export(self.payload([case]), draft=True)

        self.assertNotEqual(proc.returncode, 0)
        self.assertTrue(any("symlink" in issue.lower() for issue in result["issues"]))
        self.assertFalse(any((output / "attachments").glob("**/outside-link")))

    def test_draft_with_missing_capture_directory_does_not_copy_input_parent(self):
        capture = {
            "id": "capture-a",
            "dir": "",
            "tracePath": "trace/session.jsonl",
            "codePath": "code",
            "hash": "capture-a",
            "createdAt": 1,
        }
        pending = self.round("prompt-a", "capture-a", 1, status="pending", evaluations=[])
        case = self.case("task-a", [pending], [capture], completed=False)

        proc, result, output = self.run_export(self.payload([case]), draft=True)

        self.assertEqual(proc.returncode, 0, proc.stderr)
        self.assertTrue(any("capture directory is missing" in issue for issue in result["issues"]))
        self.assertFalse(any((output / "attachments").glob("**/input.json")))

    @unittest.skipIf(openpyxl is None, "openpyxl is only needed for workbook behavior tests")
    def test_partial_draft_evaluation_preserves_known_scores_descriptions_and_metadata(self):
        capture = self.capture("capture-a")
        evaluation = self.evaluation("partial", status="needs_evidence")
        evaluation["scores"] = [5, None, 4, None, 3]
        evaluation["descriptions"] = ["交付证据充分", "缺少约束原文", "规划有一次返工", "缺少完整轨迹", "执行完成但验证不足"]
        evaluation["missing"] = ["instruction text", "reasoning events"]
        round_item = self.round("prompt-a", "capture-a", 1, evidence_hash="partial", evaluations=[evaluation])

        proc, result, _ = self.run_export(self.payload([self.case("task-a", [round_item], [capture])]), draft=True)

        self.assertEqual(proc.returncode, 0, proc.stderr)
        sheet = openpyxl.load_workbook(result["outputPath"], data_only=False).active
        self.assertEqual([sheet.cell(2, column).value for column in (13, 15, 17, 19, 21)], [5, None, 4, None, 3])
        self.assertEqual([sheet.cell(2, column).value for column in (14, 16, 18, 20, 22)], evaluation["descriptions"])
        self.assertEqual([sheet["F2"].value, sheet["G2"].value, sheet["H2"].value, sheet["J2"].value, sheet["X2"].value, sheet["Y2"].value], [
            evaluation["environment"], "Claude Code", evaluation["harnessVersion"], evaluation["taskType"], "Reviewer", "2026-09-12T10:30:00+08:00",
        ])

    def test_unicode_identities_use_opaque_attachment_components(self):
        capture = self.capture("源捕获")
        round_item = self.round("提示-一", "源捕获", 1, session_id="会话/一")
        case = self.case("任务/甲", [round_item], [capture])

        proc, result, output = self.run_export(self.payload([case]))

        self.assertEqual(proc.returncode, 0, proc.stderr)
        manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
        attachment_root = manifest["cases"][0]["captures"][0]["attachmentRoot"]
        self.assertRegex(attachment_root, r"^attachments/task-[0-9a-f]{64}/capture-[0-9a-f]{64}$")
        self.assertNotIn("任务", attachment_root)
        self.assertEqual(manifest["rounds"][0]["sessionId"], "会话/一")
        self.assertEqual(manifest["rounds"][0]["promptId"], "提示-一")

    def test_safe_internal_symlink_is_preserved_inside_copied_capture(self):
        capture = self.capture("capture-a")
        source = Path(capture["dir"])
        os.symlink("result.txt", source / "code" / "result-link.txt")
        case = self.case("task-a", [self.round("prompt-a", "capture-a", 1)], [capture])

        proc, result, output = self.run_export(self.payload([case]))

        self.assertEqual(proc.returncode, 0, proc.stderr)
        copied = self.attachment_root(output, "task-a", "capture-a") / "code" / "result-link.txt"
        self.assertTrue(copied.is_symlink())
        self.assertEqual(copied.read_text(encoding="utf-8"), "model artifact\n")
        copied_root = self.attachment_root(output, "task-a", "capture-a").resolve()
        self.assertEqual(os.path.commonpath([copied.resolve(), copied_root]), str(copied_root))

    @unittest.skipIf(openpyxl is None, "openpyxl is only needed for workbook behavior tests")
    def test_round_without_capture_id_uses_latest_raw_trace_containing_identity(self):
        older = self.capture("older")
        latest = self.capture("latest")
        older["createdAt"] = 1
        latest["createdAt"] = 2
        Path(latest["dir"], latest["tracePath"]).write_text(
            json.dumps({"sessionId": "session-1", "promptId": "prompt-a", "prompt": "Do the work"}) + "\n",
            encoding="utf-8",
        )
        round_item = self.round("prompt-a", "", 1)
        case = self.case("task-a", [round_item], [older, latest])

        proc, result, output = self.run_export(self.payload([case]))

        self.assertEqual(proc.returncode, 0, proc.stderr)
        manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["rounds"][0]["captureId"], "")
        self.assertEqual(manifest["rounds"][0]["traceCaptureId"], "latest")
        sheet = openpyxl.load_workbook(result["outputPath"], data_only=False).active
        self.assertTrue(sheet["E2"].value.endswith("/trace/session.jsonl"))

    def test_report_distinguishes_automated_structure_from_semantic_review(self):
        capture = self.capture("capture-a")
        case = self.case("task-a", [self.round("prompt-a", "capture-a", 1)], [capture])

        proc, result, _ = self.run_export(self.payload([case]))

        self.assertEqual(proc.returncode, 0, proc.stderr)
        report = Path(result["reportPath"]).read_text(encoding="utf-8")
        self.assertIn("Automated syntax/structure: PASS", report)
        self.assertIn("Semantic evaluation: NOT CHECKED", report)

    def test_ready_evaluation_review_directory_is_required_formally_and_copied(self):
        capture = self.capture("capture-a")
        evaluation = self.evaluation("reviewed")
        review_source = Path(evaluation["reviewPath"])
        round_item = self.round("prompt-a", "capture-a", 1, evidence_hash="reviewed", evaluations=[evaluation])
        case = self.case("task-a", [round_item], [capture])

        proc, result, output = self.run_export(self.payload([case]))

        self.assertEqual(proc.returncode, 0, proc.stderr)
        manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
        review_attachment = output / manifest["rounds"][0]["reviewAttachment"]
        self.assertRegex(review_attachment.relative_to(output).as_posix(), r"^attachments/task-[0-9a-f]{64}/reviews/evaluation-[0-9a-f]{64}$")
        self.assertEqual((review_attachment / "verification" / "check.txt").read_text(encoding="utf-8"), "independent check\n")
        evidence = (output / "evidence.md").read_text(encoding="utf-8")
        self.assertIn((review_attachment / "verification" / "check.txt").relative_to(output).as_posix(), evidence)

        evaluation["reviewPath"] = str(review_source / "missing")
        blocked_proc, blocked_result, _ = self.run_export(self.payload([case]), output_name="formal-missing-review")
        self.assertNotEqual(blocked_proc.returncode, 0)
        self.assertTrue(any("review evidence directory is missing" in issue for issue in blocked_result["issues"]))
        draft_proc, draft_result, _ = self.run_export(self.payload([case]), draft=True, output_name="draft-missing-review")
        self.assertEqual(draft_proc.returncode, 0, draft_proc.stderr)
        draft_sheet = openpyxl.load_workbook(draft_result["outputPath"], data_only=False).active
        self.assertEqual([draft_sheet.cell(2, column).value for column in (13, 15, 17, 19, 21)], [5, 4, 4, 4, 4])
        self.assertTrue(any("review evidence directory is missing" in issue for issue in draft_result["issues"]))


if __name__ == "__main__":
    unittest.main()
