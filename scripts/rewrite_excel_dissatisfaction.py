#!/usr/bin/env python3
import argparse
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook

import normalize_pinru_dissatisfaction as norm


HEADER_ROW = 1
SESSION_ID_COL = "Trae Session ID"
PROMPT_COL = "User Prompt"
TASK_TYPE_COL = "任务类型"
COMPLETION_COL = "任务是否完成"
SATISFACTION_COL = "过程与产物是否满意"
DISSATISFACTION_COL = "不满意原因"


def header_map(ws) -> dict[str, int]:
    result: dict[str, int] = {}
    for cell in ws[HEADER_ROW]:
        if cell.value is None:
            continue
        result[str(cell.value).strip()] = int(cell.column)
    return result


def cell_text(ws, row: int, col: int) -> str:
    value = ws.cell(row=row, column=col).value
    return str(value).strip() if value is not None else ""


def build_record(ws, row: int, headers: dict[str, int]) -> dict:
    review_notes = cell_text(ws, row, headers[DISSATISFACTION_COL])
    return {
        "task_id": "",
        "round_number": row - HEADER_ROW,
        "session_id": cell_text(ws, row, headers[SESSION_ID_COL]),
        "session_evaluation": "",
        "task_type": cell_text(ws, row, headers[TASK_TYPE_COL]),
        "user_prompt": cell_text(ws, row, headers[PROMPT_COL]),
        "current_prompt": cell_text(ws, row, headers[PROMPT_COL]),
        "review_notes": review_notes,
        "key_locations": "",
    }


def true_last_data_row(ws, headers: dict[str, int]) -> int:
    session_col = headers[SESSION_ID_COL]
    last_row = HEADER_ROW
    for row in range(HEADER_ROW + 1, ws.max_row + 1):
        if cell_text(ws, row, session_col):
            last_row = row
    return last_row


def rewrite_workbook(input_path: Path, output_path: Path, force: bool) -> tuple[int, int, int]:
    if input_path.resolve() == output_path.resolve():
        raise RuntimeError("output must be different from input")
    shutil.copy2(input_path, output_path)
    wb = load_workbook(output_path)
    ws = wb.active
    headers = header_map(ws)
    required = [SESSION_ID_COL, PROMPT_COL, TASK_TYPE_COL, SATISFACTION_COL, DISSATISFACTION_COL]
    missing = [name for name in required if name not in headers]
    if missing:
        raise RuntimeError("missing columns: " + ", ".join(missing))

    processed = 0
    rewritten = 0
    skipped = 0
    for row in range(HEADER_ROW + 1, ws.max_row + 1):
        satisfied = cell_text(ws, row, headers[SATISFACTION_COL])
        if satisfied != "不满意":
            continue
        reason = cell_text(ws, row, headers[DISSATISFACTION_COL])
        if not reason:
            skipped += 1
            continue
        processed += 1
        if norm.is_structured_reason(reason) and not force:
            skipped += 1
            continue
        record = build_record(ws, row, headers)
        try:
            summary = norm.build_rule_summary(record)
        except Exception:
            skipped += 1
            continue
        ws.cell(row=row, column=headers[DISSATISFACTION_COL]).value = summary
        rewritten += 1

    last_row = true_last_data_row(ws, headers)
    if last_row < ws.max_row:
        ws.delete_rows(last_row + 1, ws.max_row - last_row)

    wb.save(output_path)
    return processed, rewritten, skipped


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()

    processed, rewritten, skipped = rewrite_workbook(Path(args.input), Path(args.output), args.force)
    print(args.output)
    print(f"processed={processed} rewritten={rewritten} skipped={skipped}")


if __name__ == "__main__":
    main()
